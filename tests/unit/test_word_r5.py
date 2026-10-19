"""Word R5: all source text accounted for; one target resolution; literal text."""

import io
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from bs4 import BeautifulSoup
from docx import Document
from lxml import etree

from src.privacy_service import document_converter as dc
from src.privacy_service import word_provenance as wp
from src.privacy_service import word_safety as ws
from src.privacy_service.document_converter import convert_docx_bytes
from src.privacy_service.word_provenance import (
    CONTINUATION_NOTICE,
    Provenance,
    apply_provenance,
    missing_text,
    new_provenance,
    part_target,
)

KORPUS = Path(__file__).parents[1] / "fixtures/docx/korpus"
W = ws.W
WNS = f'xmlns:w="{W}"'


def probe(name):
    return (KORPUS / name).read_bytes()


def packed(document):
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def rewrite(content, name, change):
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(content)) as source,
        zipfile.ZipFile(out, "w") as target,
    ):
        for info in source.infolist():
            data = source.read(info)
            if info.filename == name:
                data = change(data)
            target.writestr(info, data)
    return out.getvalue()


def member(content, name):
    with zipfile.ZipFile(io.BytesIO(content)) as package:
        return package.read(name)


def ledger(xml):
    """Text ledger of one body built from raw WordprocessingML paragraphs."""
    root = etree.fromstring(f"<w:document {WNS}><w:body>{xml}</w:body></w:document>")
    provenance = Provenance("wqtest_")
    wp.record_text(
        provenance, "word/document.xml", root, [], {"footnote": set(), "endnote": set()}
    )
    return provenance


# A1: a table Writer never exports is found by the text ledger, not by identity.
def test_source_table_missing_from_export_is_appended_and_named():
    document = Document()
    document.add_paragraph("QXBODYQZ")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "QXLOSTAQZ", "QXLOSTBQZ"
    prepared = ws.prepare_word(packed(document))
    soup = BeautifulSoup("<body><p>QXBODYQZ</p></body>", "html.parser")
    notices = apply_provenance(soup, soup.body, prepared.provenance)
    assert notices == [
        "Dokumenttext: 2 Absätze im Export nicht enthalten, aus der Quelle angehängt"
    ]
    text = soup.get_text(" ")
    assert (
        "[Dokumenttext: Text im Export nicht enthalten, aus dem Dokument übernommen]"
        in text
    )
    assert text.count("QXLOSTAQZ") == 1 and text.count("QXLOSTBQZ") == 1
    assert text.count("QXBODYQZ") == 1


@pytest.mark.parametrize(
    "name",
    ["c01_fussnote_tabelle_nach_absatz.docx", "c02_fussnote_tabelle_vor_absatz.docx"],
)
def test_footnote_table_is_in_the_ledger(name):
    provenance = ws.prepare_word(probe(name)).provenance
    notes = [t for t in provenance.texts if t.kind == "notes" and t.label == "Fußnoten"]
    units = [u for t in notes for p in t.paragraphs for u in p.units]
    assert [f"QXNOTE{i}QZ" for i in range(4)] == [
        u for u in units if u.startswith("QXNOTE")
    ]


def test_ledger_ignores_field_results_and_splits_at_note_marks():
    provenance = ledger(
        '<w:p><w:r><w:t>Seite </w:t></w:r><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText>PAGE</w:instrText></w:r><w:r><w:fldChar '
        'w:fldCharType="separate"/></w:r>'
        '<w:r><w:t>7</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r>'
        '<w:r><w:t> von</w:t></w:r><w:fldSimple w:instr="NUMPAGES"><w:r><w:t>9</w:t></'
        'w:r>'
        '</w:fldSimple><w:r><w:t>Satz</w:t></w:r><w:r><w:footnoteReference w:id="2"/></'
        'w:r>'
        "<w:r><w:t> weiter</w:t></w:r></w:p>"
    )
    ((paragraph,),) = [provenance.texts[0].paragraphs]
    assert paragraph.units == ["Seite", "von", "Satz", "weiter"]
    assert paragraph.text == "Seite 7 von9Satz weiter"


def test_ledger_keeps_branches_and_text_boxes_separate():
    mc = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    provenance = ledger(
        f'<w:p><w:r><w:t>Anker</w:t></w:r><mc:AlternateContent xmlns:mc="{mc}">'
        '<mc:Choice>'
        "<w:txbxContent><w:p><w:r><w:t>BOX</w:t></w:r></w:p></w:txbxContent></mc:Choice>"
        "<mc:Fallback><w:txbxContent><w:p><w:r><w:t>BOX</w:t></w:r></w:p></w:txbxContent>"
        "</mc:Fallback></mc:AlternateContent><w:r><w:t>Ende</w:t></w:r></w:p>"
    )
    paragraphs = provenance.texts[0].paragraphs
    assert [p.units for p in paragraphs] == [["Anker", "Ende"], ["BOX"], ["BOX"]]
    group = paragraphs[1].branches[0][0]
    assert [p.branches for p in paragraphs] == [(), ((group, 0),), ((group, 1),)]


def test_missing_text_reports_only_lost_paragraphs_once():
    provenance = ledger(
        "<w:p><w:r><w:t>Eins</w:t></w:r></w:p>"
        "<w:p><w:r><w:t>Zwei Drei</w:t></w:r></w:p>"
    )
    soup = BeautifulSoup("<body><p>E i n s</p></body>", "html.parser")
    ((part, paragraphs),) = missing_text(soup.body, provenance)
    assert paragraphs == ["Zwei Drei"]
    soup = BeautifulSoup("<body><p>ZweiDrei</p><p>Eins</p></body>", "html.parser")
    assert missing_text(soup.body, provenance) == []


# A2: presence is recorded before a local flatten drops the bookmarks.
def test_flattened_header_table_is_not_reported_missing():
    provenance = ws.prepare_word(probe("c03_kopf_mehrdeutig.docx")).provenance
    header = next(p for p in provenance.parts if p.kind == "header")
    table = next(t for t in provenance.tables.values() if t.part == header.name)
    assert not table.valid
    pre = provenance.prefix
    soup = BeautifulSoup(
        f'<body><div title="header"><table><tr><td><a name="{pre}{header.marker}"></a>'
        f'<a name="{pre}t{table.key}r0c0"></a>QXAMHEAD0QZ</td><td>QXAMHEAD1QZ</td></tr>'
        "<tr><td>QXAMHEAD2QZ</td><td>QXAMHEAD3QZ</td></tr></table></div>"
        "<p>QXHEADERAMBODYQZ</p></body>",
        "html.parser",
    )
    notices = apply_provenance(soup, soup.body, provenance)
    # flattened locally — and the header it came from is not appended a second time
    assert len(notices) == 1 and notices[0].startswith("Tabelle ohne Spaltenzuordnung")
    assert all(soup.get_text().count(f"QXAMHEAD{i}QZ") == 1 for i in range(4))


# A3: stored text of a continuation cell moves to its origin cell, named.
def test_continuation_cell_text_moves_to_origin():
    prepared = ws.prepare_word(probe("c04_vmerge_folgezelle_text.docx"))
    assert any("Folgezelle" in n for n in prepared.provenance.notices)
    root = etree.fromstring(member(prepared.content, "word/document.xml"))
    cells = root.findall(".//w:tc", ws.NS)
    texts = ["".join(t.text or "" for t in c.iter(f"{{{W}}}t")) for c in cells]
    origin = next(t for t in texts if "QXVMCELL0QZ" in t)
    assert CONTINUATION_NOTICE in origin and "QXCONTINUATIONQZ" in origin
    assert sum("QXCONTINUATIONQZ" in t for t in texts) == 1
    for cell in cells:
        assert cell.findall("w:p", ws.NS), "every cell keeps a paragraph"


def test_empty_continuation_cell_changes_nothing():
    prepared = ws.prepare_word(probe("h11_lange_tabelle_verbunden.docx"))
    assert not any("Folgezelle" in n for n in prepared.provenance.notices)


# A4: one resolution for every relationship target.
@pytest.mark.parametrize(
    "source,target,expected",
    [
        ("word/document.xml", "media/image1.png", "word/media/image1.png"),
        ("word/document.xml", "/word/media/image1.png", "word/media/image1.png"),
        ("word/document.xml", "media/image%201.png", "word/media/image 1.png"),
        ("word/document.xml", "../word/media/a.png", "word/media/a.png"),
        ("word/document.xml", "../../etc/passwd", None),
        ("", "/content/document.xml", "content/document.xml"),
        ("word/document.xml", "media\\a.png", "word/media/a.png"),
    ],
)
def test_part_target_resolution(source, target, expected):
    assert part_target(source, target) == expected


def test_part_target_prefers_an_existing_member():
    names = {"word/media/a%20b.png"}
    assert (
        part_target("word/document.xml", "media/a%20b.png", names)
        == "word/media/a%20b.png"
    )
    assert part_target("word/document.xml", "media/x.png", names) is None


@pytest.mark.parametrize(
    "name", ["c10_bild_paketabsolut.docx", "c11_hauptteil_verlegt_bild.docx"]
)
def test_pixel_budget_holds_for_absolute_and_relocated_targets(name, monkeypatch):
    monkeypatch.setattr(
        dc.subprocess, "run", lambda *a, **k: pytest.fail("Writer started")
    )
    prepared = ws.prepare_word(probe(name))
    assert prepared.provenance.notices == [
        "Eingebettetes Bild zu groß (30000×30000 Pixel), nicht geladen"
    ]
    with zipfile.ZipFile(io.BytesIO(prepared.content)) as package:
        rels = [n for n in package.namelist() if n.endswith(".rels")]
        assert not any(b'/image"' in package.read(n) for n in rels if "document" in n)


def test_missing_picture_member_is_not_passed_to_writer():
    raw = (
        probe("h19_pixelbombe.docx")
        if (KORPUS / "h19_pixelbombe.docx").exists()
        else probe("c10_bild_paketabsolut.docx")
    )
    raw = rewrite(
        raw,
        "word/_rels/document.xml.rels",
        lambda d: d.replace(b"media/image1.png", b"media/fehlt.png"),
    )
    prepared = ws.prepare_word(raw)
    document = member(prepared.content, "word/document.xml")
    assert ws.MISSING_IMAGE_NOTICE.encode() in document
    assert b"fehlt.png" not in member(prepared.content, "word/_rels/document.xml.rels")


def test_relocated_main_document_is_sanitized():
    prepared = ws.prepare_word(probe("c12_hauptteil_verlegt_privat.docx"))
    main = member(prepared.content, "content/document.xml")
    assert b"VERSTECKTERTEXT" not in main and b"GELOESCHTERTEXT" not in main
    assert b"KOMMENTARTEXT" not in member(prepared.content, "content/comments.xml")
    rels = member(prepared.content, "content/_rels/document.xml.rels")
    assert b"example.invalid" not in rels


# A5: ordinary text never collides with the table placeholder. Real Writer
# export: the merged table carries provenance and takes the HTML-table path.
def test_placeholder_is_redrawn_when_the_document_contains_it(monkeypatch):
    # Only the converter's draw is scripted; the provenance prefix keeps the
    # real generator, otherwise it would consume the colliding first token.
    tokens = iter(["aaaaaaaaaaaa", "bbbbbbbbbbbb"])
    drawn = []

    def token_hex(n):
        drawn.append(next(tokens))
        return drawn[-1]

    monkeypatch.setattr(dc, "secrets", SimpleNamespace(token_hex=token_hex))
    document = Document()
    document.add_paragraph("QXVORQZ WORDTABLEaaaaaaaaaaaaN0END QXNACHQZ")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "QXKOPFQZ"
    table.cell(1, 0).text = "QXLINKSQZ"
    table.cell(1, 1).text = "QXRECHTSQZ"
    markdown, _, _ = convert_docx_bytes(packed(document))
    assert "QXVORQZ WORDTABLEaaaaaaaaaaaaN0END QXNACHQZ" in markdown
    assert markdown.count("QXKOPFQZ") == 1
    assert 'colspan="2"' in markdown
    assert "WORDTABLEbbbbbbbbbbbb" not in markdown
    assert drawn == ["aaaaaaaaaaaa", "bbbbbbbbbbbb"]


# A6: collisions are checked against decoded bookmark names.
@pytest.mark.parametrize("spelling", ["wq&#x61;aaaaaaa_t1r0c0", "wqaaaaaaaa_t1r0c0"])
def test_prefix_collision_uses_decoded_names(spelling, monkeypatch):
    raw = (
        f'<w:document {WNS}><w:body><w:p><w:bookmarkStart w:id="1" w:name="{spelling}"/'
        '>'
        "</w:p></w:body></w:document>"
    ).encode()
    tokens = iter(["aaaaaaaa", "bbbbbbbb"])
    monkeypatch.setattr(wp.secrets, "token_hex", lambda n: next(tokens))
    assert new_provenance({"word/document.xml": raw}).prefix == "wqbbbbbbbb_"


# A8: text that looks like markup stays text after CommonMark rendering.
def test_literal_markup_stays_visible_text(monkeypatch):
    from tests.unit.test_docx_paragraphs import _export_html

    markdown_it = pytest.importorskip("markdown_it")
    _export_html(
        monkeypatch,
        '<body><p>QX &lt;a name="wq_x"&gt; &amp;amp; &lt;!-- c --&gt; '
        "a &lt; b QZ</p></body>",
    )
    markdown, _, _ = convert_docx_bytes(probe("spaces.docx"))
    first = markdown.split("\n\n")[0]
    html = markdown_it.MarkdownIt("commonmark").render(first)
    shown = BeautifulSoup(html, "html.parser").get_text()
    assert shown.strip() == 'QX <a name="wq_x"> &amp; <!-- c --> a < b QZ'
