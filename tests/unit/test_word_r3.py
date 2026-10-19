"""Word R3 regressions through real Writer and bounded package parsing."""

import base64
import io
import zipfile

import pytest
from bs4 import BeautifulSoup
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from src.privacy_service import word_safety as ws
from src.privacy_service.document_converter import (
    DocumentTooLargeError,
    convert_docx_bytes,
    convert_image_bytes,
)


def packed(document):
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def element(tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn(key), value)
    return node


def test_format_boundaries():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Fett ").bold = True
    p.add_run("Text danach.")
    p = doc.add_paragraph("Vorher")
    p.add_run(" kursiv").italic = True
    p.add_run(" und weiter.")
    markdown, _, _ = convert_docx_bytes(packed(doc))
    assert "**Fett** Text danach." in markdown
    assert "Vorher *kursiv* und weiter." in markdown


def test_cell_break():
    doc = Document()
    p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    p.add_run("Zeile a").add_break()
    p.add_run("Zeile b")
    markdown, _, _ = convert_docx_bytes(packed(doc))
    assert "Zeile a Zeile b" in markdown


@pytest.mark.parametrize("part", ["body", "header", "footer"])
@pytest.mark.parametrize("merged", [False, True])
def test_ordinary_word_content_in_tables(part, merged):
    doc = Document()
    doc.add_paragraph("Dokument bleibt erhalten")
    parent = doc if part == "body" else getattr(doc.sections[0], part)
    args = {} if part == "body" else {"width": doc.sections[0].page_width}
    table = parent.add_table(rows=2, cols=3, **args)
    for row in table.rows:
        for cell in row.cells:
            cell.text = "Zelle"
    if merged:
        table.cell(0, 0).merge(table.cell(0, 1))
    p = table.cell(0, 0).paragraphs[0]
    p.text = "WP"
    p.add_run()._r.append(element("w:noBreakHyphen"))
    p.add_run("1 ")
    p.add_run()._r.append(element("w:sym", **{"w:font": "Symbol", "w:char": "F0B0"}))
    p.add_run("C")
    p = table.cell(1, 2).paragraphs[0]
    p.text = ""
    field = element(
        "w:fldSimple",
        **{"w:instr": 'DATE \\@ "dd.MM.yyyy"' if part == "body" else "PAGE"},
    )
    run = element("w:r")
    text = element("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    p._p.append(field)
    markdown, _, _ = convert_docx_bytes(packed(doc))
    assert "Dokument bleibt erhalten" in markdown
    assert "WP‑1" in markdown
    assert "°C" in markdown
    if merged:
        assert BeautifulSoup(markdown, "html.parser").find("td", colspan="2")


def test_zip_budget_before_read(monkeypatch):
    doc = Document()
    doc.add_paragraph("valid")
    out = io.BytesIO(packed(doc))
    with zipfile.ZipFile(out, "a", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/media/bomb.bin", b"\0" * 2_000_000)
    monkeypatch.setattr(ws, "MAX_WORD_EXPANDED", 1_000_000)
    original = zipfile.ZipFile.read

    def guarded(self, name, *args, **kwargs):
        assert name != "word/media/bomb.bin", "Bomb inflated before budget check"
        return original(self, name, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "read", guarded)
    with pytest.raises(DocumentTooLargeError):
        ws.sanitize_docx(out.getvalue())


@pytest.mark.parametrize(
    "budget", ["MAX_WORD_ENTRIES", "MAX_WORD_XML_PART", "MAX_WORD_XML_TOTAL"]
)
def test_other_zip_budgets(monkeypatch, budget):
    monkeypatch.setattr(ws, budget, 1)
    with pytest.raises(DocumentTooLargeError):
        ws.sanitize_docx(packed(Document()))


@pytest.mark.parametrize("fmt", ["BMP", "TIFF"])
def test_raster_image_conversion(fmt):
    image = io.BytesIO()
    Image.new("RGB", (64, 48), (200, 10, 10)).save(image, format=fmt)
    image.seek(0)
    doc = Document()
    doc.add_paragraph("Bildtext bleibt")
    doc.add_picture(image)
    markdown, _, images = convert_docx_bytes(packed(doc))
    assert "Bildtext bleibt" in markdown
    assert images
    for name, data in images.items():
        assert name.endswith(".png")
        assert base64.b64decode(data).startswith(b"\x89PNG")


def test_unreadable_image_notice(monkeypatch):
    from tests.unit.test_docx_paragraphs import _export_html

    _export_html(
        monkeypatch,
        '<body><p>Text bleibt</p><img src="figure.png"/></body>',
        image=b"unrenderable metafile",
    )
    markdown, _, images = convert_docx_bytes(packed(Document()))
    assert "Text bleibt" in markdown
    assert "Eingebettetes Bild nicht verarbeitbar" in markdown
    assert not images


@pytest.mark.parametrize("kind", ["row", "cell"])
def test_tracked_structure_deletion(kind):
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    for row in table.rows:
        for cell in row.cells:
            cell.text = "sichtbar"
    if kind == "row":
        table.rows[1].cells[0].text = "GELOESCHT"
        table.rows[1]._tr.get_or_add_trPr().append(element("w:del"))
    else:
        table.cell(1, 0).text = "GELOESCHT"
        table.cell(1, 0)._tc.get_or_add_tcPr().append(element("w:cellDel"))
    with zipfile.ZipFile(io.BytesIO(ws.sanitize_docx(packed(doc)))) as archive:
        from lxml import etree

        root = etree.fromstring(archive.read("word/document.xml"))
    assert "GELOESCHT" not in "".join(root.itertext())
    rows = root.findall(".//w:tr", ws.NS)
    assert len(rows) == (2 if kind == "row" else 3)
    if kind == "cell":
        assert len(rows[1].findall("w:tc", ws.NS)) == 1


def test_empty_geometry_is_removed_locally():
    """R4: a table emptied by accepted deletions disappears, the file stays."""
    doc = Document()
    doc.add_paragraph("Bleibt")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0]._tr.get_or_add_trPr().append(element("w:del"))
    with zipfile.ZipFile(io.BytesIO(ws.sanitize_docx(packed(doc)))) as archive:
        text = archive.read("word/document.xml").decode()
    assert "<w:tbl>" not in text and "Bleibt" in text


def test_arithmetic_asterisk():
    doc = Document()
    doc.add_paragraph("5*3 = 15")
    markdown, _, _ = convert_docx_bytes(packed(doc))
    assert "5*3 = 15" in markdown


@pytest.mark.parametrize(
    "content", [b"random non Word bytes", b"%PDF-1.4", b"PK\x03\x04not a zip"]
)
def test_non_word_rejected_before_writer(content, monkeypatch):
    from src.privacy_service import document_converter as dc

    monkeypatch.setattr(
        dc.subprocess,
        "run",
        lambda *a, **kw: pytest.fail("Writer received non Word input"),
    )
    with pytest.raises(ValueError, match="Kein Word-Dokument"):
        convert_docx_bytes(content)


def test_foreign_zip_rejected():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("hello.txt", "not Word")
    with pytest.raises(ValueError, match="Kein Word-Dokument"):
        convert_docx_bytes(out.getvalue())


def test_image_docstring():
    assert "InputFormat.IMAGE" in convert_image_bytes.__doc__
    assert "from docling." not in convert_image_bytes.__doc__


def test_word_zip_limit_returns_http_413(monkeypatch):
    from fastapi.testclient import TestClient

    from src.privacy_service import app as app_mod
    from src.privacy_service.adapters import build_default_chain

    monkeypatch.setattr(ws, "MAX_WORD_EXPANDED", 1)
    monkeypatch.setattr(app_mod, "_DOCUMENT_CHAIN", build_default_chain())
    with TestClient(app_mod.app) as client:
        response = client.post(
            "/document/convert",
            files={
                "file": (
                    "budget.docx",
                    packed(Document()),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
    assert response.status_code == 413, response.text
    assert "expanded size/entry budget" in response.json()["detail"]


def test_word_package_checks_main_document():
    doc = Document()
    raw = io.BytesIO(packed(doc))
    out = io.BytesIO()
    with zipfile.ZipFile(raw) as source, zipfile.ZipFile(out, "w") as target:
        for info in source.infolist():
            data = source.read(info)
            if info.filename == "word/document.xml":
                data = b"<foreign/>"
            target.writestr(info, data)
    with pytest.raises(ValueError, match="Kein Word-Dokument"):
        ws.sanitize_docx(out.getvalue())


def test_symbol_encoding_and_unmapped_notice():
    doc = Document()
    p = doc.add_paragraph()
    for code in ("F061", "F0B0"):  # Symbol alpha and degree
        p.add_run()._r.append(element("w:sym", **{"w:font": "Symbol", "w:char": code}))
    p.add_run()._r.append(element("w:sym", **{"w:font": "Custom", "w:char": "F061"}))
    with zipfile.ZipFile(io.BytesIO(ws.sanitize_docx(packed(doc)))) as archive:
        text = archive.read("word/document.xml").decode()
    assert "α" in text and "°" in text
    assert "Symbol aus Custom: U+F061 nicht als Unicode verfügbar" in text


def test_identical_text_tables_use_part_order():
    doc = Document()
    for start in (0, 1):
        doc.add_paragraph("Getrennte Tabelle")
        table = doc.add_table(rows=1, cols=3)
        table.cell(0, start).merge(table.cell(0, start + 1))
        for cell in table.rows[0]._tr.findall(qn("w:tc")):
            text = cell.find(".//" + qn("w:t"))
            if text is None:
                run = element("w:r")
                text = element("w:t")
                run.append(text)
                cell.find(qn("w:p")).append(run)
            text.text = "gleich"
    markdown, _, _ = convert_docx_bytes(packed(doc))
    tables = BeautifulSoup(markdown, "html.parser").find_all("table")
    assert len(tables) == 2
    assert [cell.get("colspan", "1") for cell in tables[0].find_all("td")] == ["2", "1"]
    assert [cell.get("colspan", "1") for cell in tables[1].find_all("td")] == ["1", "2"]
