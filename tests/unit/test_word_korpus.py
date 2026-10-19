"""Gate against the R1–R3 failure class: one odd element must not fail a Word file.

Every probe of the independent reviews goes through the real
``/document/convert``. Only real dangers (no valid OOXML package, expanded
size budget) may answer 4xx — then without starting Writer. Everything else
must return text, with local notices where an element could not be kept.
The reviewer's sizes for the bomb probes are generated here instead of
being committed. Soll per probe: ``tests/fixtures/docx/korpus/README.md``.
"""

import binascii
import io
import shutil
import struct
import zipfile
import zlib
from pathlib import Path

import pytest
from docx import Document

KORPUS = Path(__file__).parents[1] / "fixtures/docx/korpus"
MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# Real dangers: status and reason. Nothing else may be rejected.
DANGER = {
    "bombe.docx": (413, "expanded size"),
    "h17_luegengroesse.docx": (415, "Bad CRC-32"),
    "h20a_unterbudget_unreferenziert.docx": (413, "expanded size"),
    "h20b_unterbudget_bild.docx": (413, "expanded size"),
    "riesig.docx": (413, "expanded size"),
    "falsch_pdf.docx": (415, "Kein Word-Dokument"),
    "falsch_zufall.docx": (415, "Kein Word-Dokument"),
    "falsch_zip.docx": (415, "Kein Word-Dokument"),
}

# Content the reviews named per probe; every listed string must survive.
EXPECT = {
    "voll.docx": ["Absatz EINS.", "EINGEFUEGTERTEXT", "KOPFZEILENTEXT", "FUSSZEILENTEXT",
                  "FUSSNOTENTEXT", "TEXTFELDINHALT", 'colspan="2"><p>QUERVERBUNDEN'],
    "voll.doc": ["Absatz EINS.", "QUERVERBUNDEN", "nicht gegen das Original prüfbar"],
    "spaces.docx": ["**Fett** Text danach.", "Vorher *kursiv* und weiter.", "5*3 = 15"],
    "t_sym.docx": ["20°C"],
    "t_nbh.docx": ["WP‑1"],
    "tab_feld.docx": ["|"],
    "t_pagefeld.docx": ["|"],
    "tab_mehrabsatz.docx": ["Zeile a Zeile b"],
    "tab_delrow.docx": ["| Z0 | W0 |", "| Z2 | W2 |"],
    "h01_kopf_mehrfach.docx": ["KD-QUER", "KE-QUER", "Kopfzeile (erste Seite)", "Brieftext."],
    "h01b_kopf_gerade_ohne_tabelle.docx": ["Gerade", "Kopfzeile (gerade Seiten)"],
    "h02_textfeld_tabelle.docx": ['colspan="2"><p>B-QUER', "TB-A", "TB-B"],
    "h03_sdt_zeile.docx": ['colspan="2"><p>S-QUER', "S22"],
    "h05_fussnote_endnote.docx": ["FN **FETT** und *KURSIV* Wert 5*3.", "ENDNOTENTEXT"],
    "h06_kommentar_tabelle.docx": ['colspan="2"><p>KQUER'],
    "h07_rtl_emoji.docx": ["مرحبا", "👨‍👩‍👧", "🇦🇹", "𠜎", 'colspan="2"><p>E-QUER'],
    "h10_lange_tabelle.docx": [],
    "h11_lange_tabelle_verbunden.docx": [],
    "h12_liste_ueber_tabelle.docx": ["Punkt eins", "L-QUER", "innen b", "Punkt drei"],
    "h13_makro.docm": ["Makros (VBA)", "MAKROTEXT", 'colspan="2"><p>D-QUER'],
    "h15_wingdings.docx": ["Symbol aus Wingdings"],
    "h16_gridafter.docx": ['colspan="2"><p>G-QUER', "G21"],
    "h19_pixelbombe.docx": ["[Eingebettetes Bild zu groß (30000×30000 Pixel), nicht geladen]"],
    "h19b_pixelbombe_60k.docx": ["[Eingebettetes Bild zu groß (60000×60000 Pixel), nicht geladen]"],
    "h21_excel_ole.docx": ["| EXCELZELLE | 42 |", "Eingebettetes Objekt (Excel.Sheet.12)"],
    "h21b_excel_ole_word.docx": ["| EXCELZELLE | 42 |", "Eingebettetes Objekt (Excel.Sheet.12)"],
    "h24_zwei_angrenzende_tabellen.docx": ["| Einheit | kWh/m2a |", 'colspan="2"><p>Q'],
}
FORBIDDEN = {
    "voll.docx": ["VERSTECKTERTEXT", "KOMMENTARTEXT", "GELOESCHTERTEXT", "javascript:"],
    "hidden_style.docx": ["VERSTECKTABSATZ"],
    "tab_delrow.docx": ["GELOESCHT", "ZEILEWEG", "|  |  |"],
    "tab_delrow2.docx": ["GELOESCHT", "ZEILEWEG"],
    "h06_kommentar_tabelle.docx": ["KOMMENTAR"],
    "spaces.docx": ["FettText", "Vorherkursiv"],
}


def _bomb(path, megabytes, referenced):
    """Reviewer probes mk5/mk10: zeros behind a valid Word package."""
    from PIL import Image

    image = io.BytesIO()
    Image.new("RGB", (40, 40), (9, 9, 9)).save(image, "PNG")
    doc = Document()
    doc.add_paragraph("Vor")
    doc.add_picture(io.BytesIO(image.getvalue()))
    raw = io.BytesIO()
    doc.save(raw)
    source = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    block = b"\0" * (1 << 20)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            if referenced and info.filename.startswith("word/media/"):
                with target.open(info.filename, "w", force_zip64=True) as writer:
                    writer.write(source.read(info))
                    for _ in range(megabytes):
                        writer.write(block)
            else:
                target.writestr(info.filename, source.read(info))
        if not referenced:
            with target.open("word/media/blob.bin", "w", force_zip64=True) as writer:
                for _ in range(megabytes):
                    writer.write(block)


def _lying_size(path, megabytes=1536):
    """Reviewer probe h17: the directory claims 1 KiB for 1.5 GiB of zeros."""
    compressor = zlib.compressobj(9, zlib.DEFLATED, -15)
    block = b"\0" * (1 << 20)
    crc = 0
    chunks = []
    for _ in range(megabytes):
        chunks.append(compressor.compress(block))
        crc = binascii.crc32(block, crc)
    chunks.append(compressor.flush())
    payload = b"".join(chunks)
    doc = Document()
    doc.add_paragraph("Lügengröße")
    raw = io.BytesIO()
    doc.save(raw)
    data = raw.getvalue()
    name = b"word/media/luege.bin"
    cd_offset = data.find(b"PK\x01\x02")
    end_offset = data.find(b"PK\x05\x06")
    local = struct.pack("<IHHHHHIIIHH", 0x04034B50, 20, 0, 8, 0, 0, crc, len(payload), 1024,
                        len(name), 0) + name
    central = struct.pack("<IHHHHHHIIIHHHHHII", 0x02014B50, 20, 20, 0, 8, 0, 0, crc, len(payload),
                          1024, len(name), 0, 0, 0, 0, 0, cd_offset) + name
    entries, = struct.unpack("<H", data[end_offset + 10:end_offset + 12])
    directory = data[cd_offset:end_offset] + central
    end = (data[end_offset:end_offset + 8]
           + struct.pack("<HHII", entries + 1, entries + 1, len(directory),
                         cd_offset + len(local) + len(payload))
           + data[end_offset + 20:])
    path.write_bytes(data[:cd_offset] + local + payload + directory + end)


@pytest.fixture(scope="module")
def generated(tmp_path_factory):
    root = tmp_path_factory.mktemp("word-bombs")
    _bomb(root / "bombe.docx", 1536, referenced=False)
    _bomb(root / "h20a_unterbudget_unreferenziert.docx", 950, referenced=False)
    _bomb(root / "h20b_unterbudget_bild.docx", 950, referenced=True)
    _lying_size(root / "h17_luegengroesse.docx")
    return {p.name: p for p in root.iterdir()}


@pytest.fixture(scope="module")
def client():
    from fastapi.testclient import TestClient

    from src.privacy_service import app as app_mod
    from src.privacy_service.adapters import build_default_chain

    previous = app_mod._DOCUMENT_CHAIN
    app_mod._DOCUMENT_CHAIN = build_default_chain()
    with TestClient(app_mod.app) as test_client:
        yield test_client
    app_mod._DOCUMENT_CHAIN = previous


GENERATED = ["bombe.docx", "h17_luegengroesse.docx", "h20a_unterbudget_unreferenziert.docx",
             "h20b_unterbudget_bild.docx"]
PROBES = sorted(p.name for p in KORPUS.iterdir() if p.suffix in {".docx", ".doc", ".docm", ".dotx"})


def test_korpus_is_complete():
    # 48 committed reviewer probes (R1 5, R2 20 without bombe, R3 23 without h17/h20a/h20b).
    assert len(PROBES) == 48
    assert set(EXPECT) | set(FORBIDDEN) <= set(PROBES) | set(GENERATED)
    assert not set(GENERATED) & set(PROBES)


@pytest.mark.skipif(not shutil.which("libreoffice") and not shutil.which("soffice"),
                    reason="requires LibreOffice")
@pytest.mark.parametrize("name", PROBES + GENERATED)
def test_korpus_probe(name, client, generated, monkeypatch):
    from src.privacy_service import document_converter as dc

    path = generated[name] if name in generated else KORPUS / name
    if name in DANGER:
        monkeypatch.setattr(dc.subprocess, "run",
                            lambda *a, **kw: pytest.fail(f"Writer started for {name}"))
    with path.open("rb") as handle:
        response = client.post("/document/convert", files={"file": (name, handle, MIME)})
    if name in DANGER:
        status, reason = DANGER[name]
        assert response.status_code == status, response.text[:300]
        assert reason in response.json()["detail"]
        return
    assert response.status_code == 200, response.text[:500]
    body = response.json()
    markdown = body["markdown"]
    assert markdown.strip()
    for text in EXPECT.get(name, []):
        assert text in markdown, (name, text)
    for text in FORBIDDEN.get(name, []):
        assert text not in markdown, (name, text)
