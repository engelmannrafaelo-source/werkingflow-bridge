"""Word R4: element provenance instead of order matching; problems stay local."""

import io
import zipfile
from pathlib import Path

import pytest
from bs4 import BeautifulSoup
from docx import Document
from lxml import etree

from src.privacy_service import document_converter as dc
from src.privacy_service import word_safety as ws
from src.privacy_service.document_converter import (
    DocumentTooLargeError,
    convert_docx_bytes,
    detect_format,
)
from src.privacy_service.word_provenance import (
    UNCERTAIN_TABLE_NOTICE,
    Provenance,
    SourceCell,
    SourceRow,
    SourceTable,
    apply_provenance,
)

KORPUS = Path(__file__).parents[1] / "fixtures/docx/korpus"
W = ws.W


def probe(name):
    return (KORPUS / name).read_bytes()


def tables(markdown):
    return BeautifulSoup(markdown, "html.parser").find_all("table")


def cells(table):
    return [[(td.get_text(strip=True), td.get("colspan", "1"), td.get("rowspan", "1"))
             for td in tr.find_all("td", recursive=False)]
            for tr in table.find_all("tr", recursive=False)]


def rewrite(content, name, change):
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(out, "w") as target:
        for info in source.infolist():
            data = source.read(info)
            if info.filename == name:
                data = change(data)
            target.writestr(info, data)
    return out.getvalue()


def packed(document):
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# Auflage 1: mc:Choice and mc:Fallback copies are no longer counted at all.
@pytest.mark.parametrize("order", ["textbox_after", "textbox_before"])
def test_textbox_table_before_and_after_merged_table(order):
    content = probe("h02_textfeld_tabelle.docx")
    if order == "textbox_before":
        def change(data):
            root = etree.fromstring(data)
            body = root.find("w:body", ws.NS)
            table, paragraph = body[0], body[1]
            body.remove(paragraph)
            body.insert(0, paragraph)
            assert table.tag == f"{{{W}}}tbl"
            return etree.tostring(root)
        content = rewrite(content, "word/document.xml", change)
    markdown, meta, _ = convert_docx_bytes(content)
    assert markdown.count("TB-A") == 1 and markdown.count("TB-B") == 1
    merged = next(t for t in tables(markdown) if "B-QUER" in t.get_text())
    assert cells(merged)[1] == [("B-QUER", "2", "1"), ("B12", "1", "1")]
    assert meta["word_notices"] == []
    if order == "textbox_before":
        assert markdown.index("TB-A") < markdown.index("B-QUER")


# Auflage 2: header/footer variants are attached or named, never silently lost.
def test_first_page_header_with_merged_tables():
    markdown, meta, _ = convert_docx_bytes(probe("h01_kopf_mehrfach.docx"))
    found = {t.get_text(" ", strip=True).split()[0]: t for t in tables(markdown)}
    assert set(found) == {"KD00", "KE00"}
    for key, prefix in (("KD00", "KD"), ("KE00", "KE")):
        assert cells(found[key])[1] == [(f"{prefix}-QUER", "2", "1"), (f"{prefix}12", "1", "1")]
    assert "[Kopfzeile (erste Seite): im Export nicht enthalten" in markdown
    assert meta["word_notices"] == ["Kopfzeile (erste Seite) aus der Quelle angehängt"]
    assert markdown.index("KD00") < markdown.index("KE00") < markdown.index("Brieftext")


def test_even_page_header_is_appended():
    markdown, meta, _ = convert_docx_bytes(probe("h01b_kopf_gerade_ohne_tabelle.docx"))
    assert "Gerade" in markdown
    assert "[Kopfzeile (gerade Seiten): im Export nicht enthalten" in markdown
    assert "KD-QUER" in markdown
    assert meta["word_notices"] == ["Kopfzeile (gerade Seiten) aus der Quelle angehängt"]


def test_second_section_footer_is_appended():
    doc = Document()
    doc.sections[0].footer.paragraphs[0].text = "FUSS-EINS"
    doc.add_paragraph("Teil eins")
    section = doc.add_section()
    section.footer.is_linked_to_previous = False
    section.footer.paragraphs[0].text = "FUSS-ZWEI"
    doc.add_paragraph("Teil zwei")
    markdown, meta, _ = convert_docx_bytes(packed(doc))
    assert "FUSS-EINS" in markdown and "FUSS-ZWEI" in markdown
    assert markdown.count("FUSS-ZWEI") == 1
    if "Fußzeile (Abschnitt 2)" in markdown:
        assert meta["word_notices"] == ["Fußzeile (Abschnitt 2) aus der Quelle angehängt"]


# Auflage 3: rows/cells through content controls, gridAfter, adjacent tables.
def test_row_inside_content_control():
    markdown, _, _ = convert_docx_bytes(probe("h03_sdt_zeile.docx"))
    (table,) = tables(markdown)
    assert [c[0] for c in cells(table)[1]] == ["S-QUER", "S12"]
    assert cells(table)[1][0][1] == "2"
    assert [c[0] for c in cells(table)[2]] == ["S20", "S21", "S22"]


def test_grid_after_filler_cells_are_not_content():
    markdown, _, _ = convert_docx_bytes(probe("h16_gridafter.docx"))
    (table,) = tables(markdown)
    assert cells(table)[1] == [("G-QUER", "2", "1"), ("G12", "1", "1")]
    assert cells(table)[2] == [("G20", "1", "1"), ("G21", "1", "1")]


def test_writer_joined_adjacent_tables_are_split_again():
    markdown, meta, _ = convert_docx_bytes(probe("h24_zwei_angrenzende_tabellen.docx"))
    assert "| Einheit | kWh/m2a |" in markdown
    (merged,) = tables(markdown)
    assert cells(merged)[0] == [("Q", "2", "1")]
    assert "Einheit" not in merged.get_text()
    assert meta["word_notices"] == []


def _provenance(spanned=True):
    provenance = Provenance("wqtest_")
    provenance.tables["t1"] = SourceTable("1", "word/document.xml", [
        SourceRow(0, [SourceCell("t1r0c0", 2 if spanned else 1)]),
        SourceRow(0, [SourceCell("t1r1c0", 1), SourceCell("t1r1c1", 1)]),
    ], True)
    return provenance


def _html(rows, extra=""):
    body = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in row) + "</tr>" for row in rows)
    return BeautifulSoup(f"<body><p>Davor</p><table>{body}</table>{extra}<p>Danach</p></body>",
                         "html.parser")


def test_verified_table_takes_geometry_from_source_not_writer():
    soup = _html([['<a name="wqtest_t1r0c0"></a>A'],
                  ['<a name="wqtest_t1r1c0"></a>B', '<a name="wqtest_t1r1c1"></a>C']])
    soup.find("td")["colspan"] = "1"
    soup.find_all("td")[1]["colspan"] = "2"  # Writer put the span on the wrong cell
    assert apply_provenance(soup, soup.body, _provenance()) == []
    assert cells(soup.find("table")) == [[("A", "2", "1")], [("B", "1", "1"), ("C", "1", "1")]]
    assert "wqtest_" not in str(soup)


@pytest.mark.parametrize("rows,reason", [
    ([['<a name="wqtest_t1r0c0"></a>A'], ['<a name="wqtest_t1r1c0"></a>B', 'UNMARKIERT']],
     "Zelle mit Inhalt ohne Quellzuordnung"),
    ([['<a name="wqtest_t1r0c0"></a>A'], ['<a name="wqtest_t1r1c1"></a>C']],
     "Zellfolge weicht von der Quelle ab"),
    ([['<a name="wqtest_t1r0c0"></a>A', '<a name="wqtest_t1r1c0"></a>B']],
     "eine Exportzeile mischt Quellzeilen"),
])
def test_ambiguous_table_is_local_and_visible(rows, reason):
    soup = _html(rows, extra='<table><tr><td>Andere</td><td>Tabelle</td></tr></table>')
    soup.find("td")["colspan"] = "2"
    notices = apply_provenance(soup, soup.body, _provenance())
    assert notices == [f"Tabelle ohne Spaltenzuordnung ausgegeben: {reason}"]
    text = soup.get_text(" ", strip=True)
    assert UNCERTAIN_TABLE_NOTICE in text
    assert "Davor" in text and "Danach" in text and "A" in text
    # Writer's own span is not passed on, the unaffected table stays a table.
    assert soup.find(attrs={"colspan": True}) is None
    assert len(soup.find_all("table")) == 1


def test_unmarked_writer_span_is_not_trusted():
    soup = BeautifulSoup('<body><table><tr><td colspan="2">X</td></tr><tr><td>a</td><td>b</td></tr>'
                         "</table></body>", "html.parser")
    notices = apply_provenance(soup, soup.body, Provenance("wqtest_"))
    assert notices == ["Tabelle ohne Spaltenzuordnung ausgegeben: Exporttabelle ohne Quellzuordnung"]
    assert soup.find("table") is None and "X" in soup.get_text()


def test_invalid_vertical_merge_is_local():
    doc = Document()
    doc.add_paragraph("Bleibt")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "OBEN"
    table.cell(1, 1).text = "UNTEN"
    tc = table.rows[1].cells[0]._tc
    pr = tc.get_or_add_tcPr()
    etree.SubElement(pr, f"{{{W}}}vMerge")  # continuation without origin cell
    markdown, meta, _ = convert_docx_bytes(packed(doc))
    assert "Bleibt" in markdown and "OBEN" in markdown and "UNTEN" in markdown
    assert meta["word_notices"] == [
        "Tabelle ohne Spaltenzuordnung ausgegeben: vertikale Verbindung ohne Ursprungszelle"]


def test_empty_geometry_after_tracked_deletion_is_local():
    doc = Document()
    doc.add_paragraph("Text bleibt")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0]._tr.get_or_add_trPr().append(etree.Element(f"{{{W}}}del"))
    with zipfile.ZipFile(io.BytesIO(ws.sanitize_docx(packed(doc)))) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    assert root.find(".//w:tbl", ws.NS) is None
    markdown, _, _ = convert_docx_bytes(packed(doc))
    assert "Text bleibt" in markdown


# Auflage 4: budget measured against what Writer keeps in memory.
def _padded_picture(megabytes):
    from PIL import Image

    image = io.BytesIO()
    Image.new("RGB", (8, 8), (9, 9, 9)).save(image, "PNG")
    doc = Document()
    doc.add_paragraph("Vor")
    doc.add_picture(io.BytesIO(image.getvalue()))
    source = zipfile.ZipFile(io.BytesIO(packed(doc)))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            if info.filename.startswith("word/media/"):
                with target.open(info.filename, "w", force_zip64=True) as writer:
                    writer.write(source.read(info))
                    for _ in range(megabytes):
                        writer.write(b"\0" * (1 << 20))
            else:
                target.writestr(info.filename, source.read(info))
    return out.getvalue()


def test_inflating_referenced_part_rejected_before_writer(monkeypatch):
    content = _padded_picture(12)  # 12 MiB from ~12 KiB: far beyond any real image
    assert len(content) < 100_000
    monkeypatch.setattr(dc.subprocess, "run", lambda *a, **kw: pytest.fail("Writer started"))
    with pytest.raises(DocumentTooLargeError, match="expanded size"):
        convert_docx_bytes(content)


def test_total_budget_keeps_writer_under_worker_limit():
    assert ws.MAX_WORD_EXPANDED <= 256 * 1024 * 1024
    assert ws.MAX_WORD_RATIO_FREE <= 10 * 1024 * 1024


def test_small_inflation_below_threshold_converts():
    markdown, _, _ = convert_docx_bytes(_padded_picture(2))
    assert "Vor" in markdown


# Auflage 5: Word detection as wide as OOXML allows; macros removed loudly.
def test_default_content_type_template_and_macro_documents():
    for name in ("h22_default_contenttype.docx", "h14_vorlage.dotx"):
        markdown, _, _ = convert_docx_bytes(probe(name))
        assert markdown.strip()
    content = probe("h13_makro.docm")
    with zipfile.ZipFile(io.BytesIO(ws.sanitize_docx(content))) as archive:
        assert "word/vbaProject.bin" not in archive.namelist()
        assert b"vbaProject" not in archive.read("word/_rels/document.xml.rels")
    markdown, meta, _ = convert_docx_bytes(content)
    assert markdown.startswith(ws.MACRO_NOTICE)
    assert "MAKROTEXT" in markdown
    assert "Makros (VBA) entfernt, nicht ausgeführt" in meta["word_notices"]


@pytest.mark.parametrize("filename,mime", [
    ("a.docm", None), ("a.dotx", None), ("a.dotm", None),
    ("x", "application/vnd.ms-word.document.macroEnabled.12"),
    ("x", "application/vnd.openxmlformats-officedocument.wordprocessingml.template"),
    ("x", "application/vnd.ms-word.template.macroEnabled.12"),
])
def test_word_variants_reach_word_adapter(filename, mime):
    assert detect_format(filename, mime) == "docx"


def test_foreign_main_content_type_rejected():
    content = rewrite(probe("h22_default_contenttype.docx"), "[Content_Types].xml",
                      lambda d: d.replace(b"wordprocessingml.document.main+xml", b"spreadsheetml.sheet.main+xml"))
    with pytest.raises(ValueError, match="Kein Word-Dokument"):
        ws.sanitize_docx(content)


# Auflage 6: embedded objects are never silently dropped.
def test_embedded_excel_keeps_cell_values():
    markdown, meta, _ = convert_docx_bytes(probe("h21b_excel_ole_word.docx"))
    assert "[Eingebettetes Objekt (Excel.Sheet.12): Tabelleninhalt des aktiven Blatts übernommen" in markdown
    assert "| EXCELZELLE | 42 |" in markdown
    assert markdown.index("Vor Excel") < markdown.index("EXCELZELLE") < markdown.index("Nach Excel")
    assert len(meta["word_notices"]) == 1


def test_unreadable_embedded_object_gets_notice():
    content = rewrite(probe("h21b_excel_ole_word.docx"), "word/embeddings/Microsoft_Excel_Worksheet.xlsx",
                      lambda d: b"kein Arbeitsblatt")
    markdown, meta, _ = convert_docx_bytes(content)
    assert "[Eingebettetes Objekt (Excel.Sheet.12) nicht übernommen]" in markdown
    assert "Vor Excel" in markdown and "Nach Excel" in markdown
    assert meta["word_notices"] == ["Eingebettetes Objekt (Excel.Sheet.12) nicht übernommen"]


def test_grid_before_offsets_vertical_merge():
    """R3 rejected gridBefore; now it shifts the column like Word does."""
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    for cell, text in zip(table.rows[0].cells, ["X", "A", "B"]):
        cell.text = text
    first = table.rows[0]._tr.findall(f"{{{W}}}tc")
    etree.SubElement(first[1].get_or_add_tcPr(), f"{{{W}}}vMerge").set(f"{{{W}}}val", "restart")
    second = table.rows[1]._tr
    cont, last, spare = second.findall(f"{{{W}}}tc")
    second.remove(spare)
    etree.SubElement(cont.get_or_add_tcPr(), f"{{{W}}}vMerge")
    last.find(f".//{{{W}}}p").append(etree.fromstring(
        f'<w:r xmlns:w="{W}"><w:t>C</w:t></w:r>'))
    tr_pr = second.get_or_add_trPr()
    etree.SubElement(tr_pr, f"{{{W}}}gridBefore").set(f"{{{W}}}val", "1")
    markdown, meta, _ = convert_docx_bytes(packed(doc))
    (html,) = tables(markdown)
    assert cells(html) == [[("X", "1", "1"), ("A", "1", "2"), ("B", "1", "1")],
                           [("", "1", "1"), ("C", "1", "1")]]
    assert meta["word_notices"] == []


def test_row_emptied_by_cell_deletions_is_removed_locally():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    for row, prefix in zip(table.rows, ("BLEIBT", "WEG")):
        for index, cell in enumerate(row.cells):
            cell.text = f"{prefix}{index}"
    for tc in table.rows[1]._tr.findall(f"{{{W}}}tc"):
        etree.SubElement(tc.get_or_add_tcPr(), f"{{{W}}}cellDel")
    with zipfile.ZipFile(io.BytesIO(ws.sanitize_docx(packed(doc)))) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    assert len(root.findall(".//w:tr", ws.NS)) == 1
    markdown, _, _ = convert_docx_bytes(packed(doc))
    assert "| BLEIBT0 | BLEIBT1 |" in markdown and "WEG" not in markdown


def _pixel_document(*edges):
    from PIL import Image

    doc = Document()
    doc.add_paragraph("Vor den Bildern")
    for shade, edge in enumerate(edges):
        image = io.BytesIO()
        Image.new("L", (edge, edge), shade * 40).save(image, "PNG")
        doc.add_picture(io.BytesIO(image.getvalue()))
    doc.add_paragraph("Nach den Bildern")
    return packed(doc)


def test_pixel_bomb_is_a_local_notice_before_writer():
    """h19: 30000^2 px from a few KB drove Writer to 1.07 GB."""
    content = probe("h19_pixelbombe.docx")
    with zipfile.ZipFile(io.BytesIO(ws.sanitize_docx(content))) as archive:
        xml = archive.read("word/document.xml").decode()
    assert "Eingebettetes Bild zu groß (30000×30000 Pixel), nicht geladen" in xml
    markdown, meta, images = convert_docx_bytes(content)
    assert "[Eingebettetes Bild zu groß (30000×30000 Pixel), nicht geladen]" in markdown
    assert images == {}
    assert meta["word_notices"] == ["Eingebettetes Bild zu groß (30000×30000 Pixel), nicht geladen"]


def test_total_picture_budget_keeps_first_pictures(monkeypatch):
    monkeypatch.setattr(ws, "MAX_WORD_IMAGE_PIXELS", 2_000_000)
    monkeypatch.setattr(ws, "MAX_WORD_IMAGE_PIXELS_TOTAL", 3_000_000)
    markdown, meta, images = convert_docx_bytes(_pixel_document(1000, 1000, 1200, 1500, 900))
    assert len(images) == 3  # 1000², 1000², 900²: 2.81 M px within the total
    assert markdown.count("Bildbudget des Dokuments erschöpft (1200×1200 Pixel)") == 1
    assert markdown.count("Eingebettetes Bild zu groß (1500×1500 Pixel)") == 1
    assert "Vor den Bildern" in markdown and "Nach den Bildern" in markdown
    assert len(meta["word_notices"]) == 2
